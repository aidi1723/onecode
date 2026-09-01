from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = Path("output/doc/OneCode_Agent_Kernel_Blog_Design_CN.docx")
NAVY = "17324D"
BLUE = "2F6B91"
PALE_BLUE = "EAF2F7"
PALE_GRAY = "F4F6F8"
TEXT = RGBColor(37, 45, 52)
MUTED = RGBColor(92, 106, 116)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=160, start=180, bottom=160, end=180):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, size=None, bold=None, color=None, name="Microsoft YaHei"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])
    set_run_font(run, 9, color=MUTED)


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    set_run_font(run, 15 if level == 1 else 12, True, RGBColor.from_string(NAVY))
    return paragraph


def add_body(doc, text, bold_prefix=None):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.line_spacing = 1.28
    paragraph.paragraph_format.space_after = Pt(4)
    if bold_prefix and text.startswith(bold_prefix):
        first = paragraph.add_run(bold_prefix)
        set_run_font(first, 10.2, True, TEXT)
        rest = paragraph.add_run(text[len(bold_prefix):])
        set_run_font(rest, 10.2, color=TEXT)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, 10.2, color=TEXT)
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.left_indent = Cm(0.55)
    paragraph.paragraph_format.first_line_indent = Cm(-0.25)
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(text)
    set_run_font(run, 10, color=TEXT)
    return paragraph


def add_equation(doc, text, note):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(text)
    set_run_font(run, 13, color=RGBColor.from_string(BLUE), name="Cambria Math")
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(6)
    cap_run = caption.add_run(note)
    set_run_font(cap_run, 9, color=MUTED)


doc = Document()
doc.core_properties.title = "OneCode：我们如何用离散数学构建一个安全、完全可控的 Agent 内核"
doc.core_properties.subject = "OneCode Agent 内核博客写作设计稿"
doc.core_properties.author = "OneCode"

section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(1.8)
section.bottom_margin = Cm(1.6)
section.left_margin = Cm(2.35)
section.right_margin = Cm(2.35)
section.header_distance = Cm(0.8)
section.footer_distance = Cm(0.8)

normal = doc.styles["Normal"]
normal.font.name = "Microsoft YaHei"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(10.5)
normal.font.color.rgb = TEXT

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
header_run = header.add_run("ONECODE  |  AGENT KERNEL")
set_run_font(header_run, 8.5, True, MUTED)
add_page_number(section.footer.paragraphs[0])

eyebrow = doc.add_paragraph()
eyebrow.paragraph_format.space_after = Pt(10)
eyebrow_run = eyebrow.add_run("TECHNICAL BLOG DESIGN  /  2026.07.12")
set_run_font(eyebrow_run, 9, True, RGBColor.from_string(BLUE))

title = doc.add_paragraph()
title.paragraph_format.space_after = Pt(8)
title_run = title.add_run("OneCode：我们如何用离散数学\n构建一个安全、完全可控的 Agent 内核")
set_run_font(title_run, 24, True, RGBColor.from_string(NAVY))

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(18)
subtitle_run = subtitle.add_run("博客写作设计稿")
set_run_font(subtitle_run, 11, color=MUTED)

table = doc.add_table(rows=1, cols=1)
table.autofit = True
cell = table.cell(0, 0)
set_cell_shading(cell, PALE_BLUE)
set_cell_margins(cell, 210, 250, 210, 250)
cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
claim = cell.paragraphs[0]
claim.paragraph_format.line_spacing = 1.35
claim_run = claim.add_run("大模型可以生成内容、计划和工具调用建议，但不应拥有最终执行权。OneCode 在模型输出与真实行动之间建立确定性内核，将不确定的模型输出转化为有限、受约束、可验证、可追溯、可恢复的执行过程。")
set_run_font(claim_run, 11, True, RGBColor.from_string(NAVY))

add_heading(doc, "一、写作目标")
add_body(doc, "这是一篇介绍 OneCode 的中文技术博客，主要面向 Agent 开发者和技术负责人，同时让更广泛的 AI 从业者也能读懂。")
add_body(doc, "文章不是通用 Agent 教程、产品说明书或技术白皮书，而是一篇关于 OneCode 的技术宣言：清楚说明我们做了什么、为什么这样做、具备哪些优势，同时不披露构成技术壁垒的核心公式、参数和决策规则。")

add_heading(doc, "二、核心价值")
values = doc.add_table(rows=1, cols=4)
values.autofit = True
for idx, value in enumerate(("安全", "可控", "可验证", "可恢复")):
    c = values.cell(0, idx)
    set_cell_shading(c, NAVY if idx == 0 else PALE_GRAY)
    set_cell_margins(c, 150, 100, 150, 100)
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(value)
    set_run_font(r, 11, True, RGBColor(255, 255, 255) if idx == 0 else RGBColor.from_string(NAVY))
add_body(doc, "自主知识产权、独立实现、模型无关、本地优先、证据防篡改和工程验证，是支撑这四个核心价值的项目优势。")

add_heading(doc, "三、可公开的数学表达")
add_body(doc, "文章只使用解释设计思想所必需的抽象公式，不与真实状态位、映射关系、转移条件或内部决策表建立对应。")
add_equation(doc, "X = {0, 1}ⁿ", "有限离散状态空间")
add_equation(doc, "φ : E → X", "将输入、环境和运行结果等证据投影为受控状态")
add_equation(doc, "T : X × I → X", "状态根据有效输入进行确定性转移")
add_equation(doc, "D : X → A", "内核只能从有限、受约束的行动集合中作出决定")
add_body(doc, "全文只使用现代工程和离散数学语言，不讨论传统文化、玄学或内部规则体系的历史来源，统一表述为“OneCode 自主研发的形式化规则体系”。")

add_heading(doc, "四、正文结构")
sections = [
    ("为什么我们要做 OneCode", "模型生成能力不等于安全可靠的执行能力。OneCode 的目标不是最大化模型的自由，而是控制模型是否可以行动，以及行动如何影响真实环境。"),
    ("OneCode 是什么", "OneCode 是我们从底层独立设计和实现的安全、完全可控的 Agent 内核，拥有自主知识产权和完整的技术演进控制权。它不是现有 Agent 框架的二次封装。"),
    ("安全与可控如何实现", "模型输出、Prompt、Skill、项目规则和外部建议都只是候选信息或证据，不拥有执行权。未知、异常或相互冲突的证据不会获得猜测性授权。"),
    ("一次任务如何通过内核", "用“输入 -> 证据 -> 状态 -> 安全裁决 -> 受控行动 -> 验证 -> 证据闭环 -> 交付”展示完整链路。"),
    ("OneCode 的核心优势", "集中说明安全内生、确定性控制、独立验证完成、证据防篡改、确定性恢复、可重放、可审计、模型无关、本地优先和独立实现。"),
    ("我们真正构建的是什么", "OneCode 不是聊天界面，也不是工具调用包装层。它是位于模型与真实行动之间的可信执行内核。"),
]
for idx, (name, body) in enumerate(sections, 1):
    add_body(doc, f"{idx}. {name}", bold_prefix=f"{idx}. {name}")
    add_body(doc, body)

add_heading(doc, "五、任务闭环与完成条件")
flow = doc.add_table(rows=1, cols=1)
flow_cell = flow.cell(0, 0)
set_cell_shading(flow_cell, PALE_GRAY)
set_cell_margins(flow_cell, 180, 220, 180, 220)
flow_p = flow_cell.paragraphs[0]
flow_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
flow_r = flow_p.add_run("输入  ->  证据  ->  状态  ->  安全裁决  ->  受控行动\n验证  ->  证据闭环  ->  交付")
set_run_font(flow_r, 11, True, RGBColor.from_string(NAVY))
add_body(doc, "模型不能自行宣布任务完成。任务进入可交付状态，必须同时满足以下条件：")
for item in ("状态有效", "物理结果符合约束", "验证通过", "关键证据完整"):
    add_bullet(doc, item)

add_heading(doc, "六、OneCode 核心优势")
advantages = (
    "安全内生与默认拒绝",
    "确定性、有限且受约束的执行权",
    "模型不能自行证明任务完成",
    "追加式、可交叉验证的防篡改证据",
    "确定性恢复、幂等重试和冲突识别",
    "可重现、可重放、可追溯、可审计",
    "模型无关与本地数据控制",
    "轻量核心，不依赖现成 Agent 框架",
    "单一权威链，避免组件越权和策略漂移",
    "自主设计、自主实现和自主技术演进",
    "完整的自动化验证与数学审计能力",
)
for item in advantages:
    add_bullet(doc, item)

add_heading(doc, "七、保密边界")
boundary = doc.add_table(rows=1, cols=2)
boundary.autofit = True
public_cell, private_cell = boundary.rows[0].cells
for cell_, fill in ((public_cell, PALE_BLUE), (private_cell, PALE_GRAY)):
    set_cell_shading(cell_, fill)
    set_cell_margins(cell_, 180, 200, 180, 200)

p = public_cell.paragraphs[0]
r = p.add_run("可以公开")
set_run_font(r, 11, True, RGBColor.from_string(BLUE))
for text in (
    "有限状态和确定性控制的抽象思想",
    "Agent 任务的外部生命周期",
    "安全、证据、验证和恢复原则",
    "已验证的项目能力和结果",
    "项目的独立实现与自主知识产权",
):
    para = public_cell.add_paragraph(style="List Bullet")
    run = para.add_run(text)
    set_run_font(run, 9.5, color=TEXT)

p = private_cell.paragraphs[0]
r = p.add_run("不能公开")
set_run_font(r, 11, True, RGBColor.from_string(NAVY))
for text in (
    "真实证据到状态的映射",
    "内部状态位含义",
    "状态转移表和决策矩阵",
    "权重、阈值、优先级和平衡参数",
    "冲突消解、恢复、调制和调度规则",
    "足以复刻内核的伪代码与实现细节",
):
    para = private_cell.add_paragraph(style="List Bullet")
    run = para.add_run(text)
    set_run_font(run, 9.5, color=TEXT)

add_heading(doc, "八、表述与准确性边界")
for item in (
    "使用清晰、克制、简洁的中文，少用术语，不做冗长推导，不为篇幅补充内容。",
    "将 OneCode 表述为确定性、本地优先的 Agent 内核，不表述为全能型自主助手。",
    "客观表述自主知识产权和实现控制权，不把 GPL v3 开源误解为放弃著作权。",
    "区分仓库已验证能力与未来部署能力，不把本地确定性基准误述为在线模型的原始幻觉率。",
    "没有固定字数。约 2000 字只是自然方向，以完整、准确表达 OneCode 为准。",
):
    add_bullet(doc, item)

closing = doc.add_table(rows=1, cols=1)
closing_cell = closing.cell(0, 0)
set_cell_shading(closing_cell, NAVY)
set_cell_margins(closing_cell, 220, 260, 220, 260)
closing_p = closing_cell.paragraphs[0]
closing_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
closing_r = closing_p.add_run("OneCode 不是另一个会调用工具的聊天机器人，\n而是位于模型与真实行动之间、负责控制执行权的可信 Agent 内核。")
set_run_font(closing_r, 11.5, True, RGBColor(255, 255, 255))

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
